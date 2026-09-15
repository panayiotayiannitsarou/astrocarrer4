import io
import re
from pathlib import Path

import streamlit as st
from docx import Document
from docx.oxml.ns import qn


# Το core/ είναι πλέον υποφάκελος του repo (χρησιμοποιείται και από τα τρία
# entry-point scripts: Home.py, AstroCheck_Analysis.py, AstroCheck_Career.py),
# οπότε η ρίζα του repo είναι ένα επίπεδο πάνω από αυτό το αρχείο, όχι το ίδιο
# επίπεδο όπως παλιά που το reference_loader.py ζούσε δίπλα στο app.py.
REPO_ROOT = Path(__file__).resolve().parent.parent
REFERENCE_DIR = REPO_ROOT / "references"
DEFAULT_INSTRUCTIONS = REFERENCE_DIR / "Odigies_v5.docx"
DEFAULT_STYLE = REFERENCE_DIR / "Elena_style_guide_v2.docx"
ROOT_INSTRUCTIONS = REPO_ROOT / "Odigies_v5.docx"
ROOT_STYLE = REPO_ROOT / "Elena_style_guide_v2.docx"
COMMON_ORIENTATION = REFERENCE_DIR / "Desmeftiki_Entoli_Epaggelmatikou_Prosanatolismou_Koini_v10.docx"
SHORT_ADULT_EXAMPLE = REFERENCE_DIR / "Protypo_Syntomis_Ekdosis_Enilikou.docx"
SHORT_TEEN_EXAMPLE = REFERENCE_DIR / "Protypo_Syntomis_Ekdosis_Paidiou_Efivou.docx"
_ROOT_COMMON_ORIENTATION = REPO_ROOT / "Desmeftiki_Entoli_Epaggelmatikou_Prosanatolismou_Koini_v10.docx"
_ROOT_SHORT_ADULT = REPO_ROOT / "Protypo_Syntomis_Ekdosis_Enilikou.docx"
_ROOT_SHORT_TEEN = REPO_ROOT / "Protypo_Syntomis_Ekdosis_Paidiou_Efivou.docx"


def docx_text(source) -> str:
    """Extract paragraphs and tables from a DOCX path or uploaded bytes."""
    if isinstance(source, (bytes, bytearray)):
        source = io.BytesIO(source)
    document = Document(source)
    blocks = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            line = " | ".join(cell.text.strip() for cell in row.cells)
            if line.strip(" |"):
                blocks.append(line)
    return "\n".join(blocks)


def simple_docx_format_issues(source) -> list[str]:
    """Εντοπίζει χρωματικές επισημάνσεις που απαγορεύονται στο καθαρό Word."""
    if isinstance(source, (bytes, bytearray)):
        source = io.BytesIO(source)
    document = Document(source)
    found = set()
    for paragraph in document.paragraphs:
        ppr = paragraph._p.pPr
        if ppr is not None:
            shd = ppr.find(qn("w:shd"))
            if shd is not None and shd.get(qn("w:fill"), "auto").lower() not in ("auto", "ffffff", "clear", "nil"):
                found.add("σκίαση παραγράφου")
        for run in paragraph.runs:
            if run.font.highlight_color is not None:
                found.add("highlight")
            rpr = run._r.rPr
            if rpr is not None:
                shd = rpr.find(qn("w:shd"))
                if shd is not None and shd.get(qn("w:fill"), "auto").lower() not in ("auto", "ffffff", "clear", "nil"):
                    found.add("χρωματιστό φόντο κειμένου")
            if run.font.color.rgb is not None and str(run.font.color.rgb).upper() not in ("000000", "FFFFFF"):
                found.add("έγχρωμο κείμενο")
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                tcpr = cell._tc.tcPr
                shd = tcpr.find(qn("w:shd")) if tcpr is not None else None
                if shd is not None and shd.get(qn("w:fill"), "auto").lower() not in ("auto", "ffffff", "clear", "nil"):
                    found.add("σκίαση πίνακα")
    return sorted(found)


@st.cache_data(show_spinner=False)
def load_default_references() -> tuple[str, str]:
    # @st.cache_data: το Streamlit ξανατρέχει ολόκληρο το script σε κάθε
    # interaction, οπότε χωρίς caching αυτά τα (συχνά εκτενή) .docx
    # ξαναδιαβάζονταν και ξαναπαρσάρονταν από τον δίσκο σε κάθε κλικ.
    #
    # Accept both repository layouts: a dedicated references/ folder or the
    # two DOCX files beside app.py.  This makes GitHub web uploads simpler.
    instructions = DEFAULT_INSTRUCTIONS if DEFAULT_INSTRUCTIONS.exists() else ROOT_INSTRUCTIONS
    style = DEFAULT_STYLE if DEFAULT_STYLE.exists() else ROOT_STYLE
    if not instructions.exists() or not style.exists():
        # Πριν έγραφε "...v4...", ενώ το πραγματικό αρχείο είναι Odigies_v5.docx
        # (και το app.py το παρουσιάζει ως "Ενσωματωμένες οδηγίες v5.3") --
        # ένα μήνυμα σφάλματος έπρεπε τουλάχιστον να συμφωνεί με το filename.
        raise FileNotFoundError("Λείπουν οι ενσωματωμένες οδηγίες v5 ή το πρότυπο ύφους.")
    return docx_text(instructions), docx_text(style)



# Η εφαρμογή έχει κλειδώσει μόνιμα την "Απλή και πρακτική" παρουσίαση (βλ. app.py,
# η "Αναλυτική" αφαιρέθηκε). Πριν, το load_orientation_command αγνοούσε το service
# και επέστρεφε ΟΛΟΚΛΗΡΟ το ~7.000 λέξεων έγγραφο των κανόνων -- και τα δύο modes
# (0Γ ενήλικα ΚΑΙ 0Δ παιδιού) μαζί, ΚΑΙ όλες τις ενότητες της "Αναλυτικής" έκδοσης
# που δεν χρησιμοποιούνται ποτέ πια. Αυτό έθαβε το μοναδικό ουσιαστικό σήμα (ποιο
# mode ισχύει) μέσα σε χιλιάδες λέξεις άσχετου/αντικρουόμενου κειμένου -- πιθανή
# αιτία όταν μια παραγωγή αγνόησε την "Απλή" λειτουργία και έβγαλε πλήρη αναλυτική
# έκδοση 7.000+ λέξεων.
#
# Ενότητες που αφορούν αποκλειστικά την "Αναλυτική" έκδοση (ή το άλλο mode) και
# ρητά υπερισχύονται/παραλείπονται από τους Κανόνες 0Γ/0Δ -- ασφαλές να μην
# σταλούν καθόλου, αφού δεν εφαρμόζονται ποτέ σε αυτή την εφαρμογή:
_ANALYTICAL_ONLY_HEADINGS = (
    "0Α.",   # πίνακες ταλέντων -- το 0Δ λέει ρητά "Μην χρησιμοποιήσεις πίνακες"
    "10Α.",  # πλήρες Πλαίσιο Εκπαιδευτικού Συστήματος -- το 0Γ/0Δ ορίζουν τη δική τους, σύντομη εκδοχή
    "11.",   # πλήρεις επαγγελματικές οικογένειες -- υπερισχύεται από το σύντομο σχήμα του 0Γ/0Δ
    "12.",   # πλήρης ενσωμάτωση επαγγελμάτων -- ίδιος λόγος
    "13.",   # περιβάλλον εργασίας -- ρητά παραλείπεται στη σύντομη έκδοση
    "13Α.",  # δυνατά σημεία/εμπόδια -- ρητά παραλείπονται στη σύντομη έκδοση
    "14.",   # σχέδιο διερεύνησης 8-12 εβδομάδων -- ρητά παραλείπεται
    "15.",   # οδηγίες προς γονείς -- ρητά παραλείπεται
    "16.",   # πλήρης υποχρεωτική δομή -- υπερισχύεται από τη δομή του 0Γ/0Δ
    "18.",   # Συμβολική Κατεύθυνση Εξέλιξης -- ρητά παραλείπεται στη σύντομη έκδοση
)


def _filter_orientation_sections(text: str, service: str) -> str:
    lines = text.split("\n")
    heading_re = re.compile(r"^(0[Α-Ω]?\.|[0-9]{1,2}[Α-Ω]?\.)\s")

    other_mode_heading = (
        "0Γ." if service == "Παιδί/έφηβος" else "0Δ."
    )
    excluded_prefixes = _ANALYTICAL_ONLY_HEADINGS + (other_mode_heading,)

    headings = [i for i, l in enumerate(lines) if heading_re.match(l)]
    drop_ranges = []
    for pos, idx in enumerate(headings):
        heading_text = lines[idx]
        if heading_text.startswith(excluded_prefixes):
            end = headings[pos + 1] if pos + 1 < len(headings) else len(lines)
            drop_ranges.append((idx, end))

    keep = [True] * len(lines)
    for start, end in drop_ranges:
        for i in range(start, end):
            keep[i] = False

    return "\n".join(l for l, k in zip(lines, keep) if k)


def load_orientation_command(service: str) -> str:
    path = COMMON_ORIENTATION if COMMON_ORIENTATION.exists() else _ROOT_COMMON_ORIENTATION
    if not path.exists():
        raise FileNotFoundError(f"Λείπει η κοινή εντολή προσανατολισμού: {COMMON_ORIENTATION.name}")
    full_text = docx_text(path)
    return _filter_orientation_sections(full_text, service)


def load_short_adult_example() -> str:
    """Ανώνυμο πρότυπο μορφής· δεν αποτελεί πηγή προσωπικών συμπερασμάτων."""
    path = SHORT_ADULT_EXAMPLE if SHORT_ADULT_EXAMPLE.exists() else _ROOT_SHORT_ADULT
    if not path.exists():
        raise FileNotFoundError(f"Λείπει το πρότυπο σύντομης έκδοσης: {SHORT_ADULT_EXAMPLE.name}")
    return docx_text(path)


def load_short_teen_example() -> str:
    """Ανώνυμο πρότυπο απλής έκδοσης παιδιού/εφήβου."""
    path = SHORT_TEEN_EXAMPLE if SHORT_TEEN_EXAMPLE.exists() else _ROOT_SHORT_TEEN
    if not path.exists():
        raise FileNotFoundError(f"Λείπει το πρότυπο παιδιού/εφήβου: {SHORT_TEEN_EXAMPLE.name}")
    return docx_text(path)
